"""process_document builds a langchain Blob from the FileUploadDTO's path
rather than reading the file into memory itself (spec T4): the DTO's
own content is never held as an in-process `bytes` value at any point
in this module. DOCX/XLSX/PPTX are loaded directly from that path (all
three libraries accept a path natively) instead of the previous
`io.BytesIO(blob.as_bytes())`, which fully materialised the file in
memory a second time even though it was already spooled to disk.
"""

from __future__ import annotations

import hashlib
import io
import json
import tempfile
from pathlib import Path

import pytest
from error_contract import ApplicationError
from langchain_core.documents.base import Blob

from src.documents.processor import (
    MAX_FILE_SIZE_BYTES,
    process_document,
)
from src.documents.schemas import FileUploadDTO


def _dto(
    content: bytes, *, filename: str, content_type: str, size: int | None = None
) -> FileUploadDTO:
    fd, raw_path = tempfile.mkstemp(prefix="test-doc-", suffix=Path(filename).suffix)
    path = Path(raw_path)
    path.write_bytes(content)
    import os

    os.close(fd)
    return FileUploadDTO(
        filename=filename,
        content_type=content_type,
        size=size if size is not None else len(content),
        path=path,
        content_hash=hashlib.sha256(content).hexdigest(),
    )


@pytest.fixture
def cleanup_dtos():
    created: list[FileUploadDTO] = []
    yield created
    for dto in created:
        dto.cleanup()


@pytest.mark.asyncio
async def test_parses_a_small_text_file_via_its_spooled_path(cleanup_dtos):
    dto = _dto(
        b"Hello, world!\n\nSecond paragraph here.",
        filename="a.txt",
        content_type="text/plain",
    )
    cleanup_dtos.append(dto)

    docs = await process_document(dto)

    assert any("Hello, world!" in d.page_content for d in docs)


@pytest.mark.asyncio
async def test_parses_a_csv_file_via_its_spooled_path(cleanup_dtos):
    csv_bytes = b"name,age\nAlice,30\nBob,25\n"
    dto = _dto(csv_bytes, filename="a.csv", content_type="text/csv")
    cleanup_dtos.append(dto)

    docs = await process_document(dto)

    assert len(docs) == 2
    assert "name: Alice" in docs[0].page_content
    assert "age: 30" in docs[0].page_content


@pytest.mark.asyncio
async def test_parses_a_json_file_via_its_spooled_path(cleanup_dtos):
    payload = json.dumps({"title": "Report", "author": "Ada"}).encode()
    dto = _dto(payload, filename="a.json", content_type="application/json")
    cleanup_dtos.append(dto)

    docs = await process_document(dto)

    assert any("title: Report" in d.page_content for d in docs)


@pytest.mark.asyncio
async def test_rejects_a_file_whose_declared_size_exceeds_the_limit_without_reading_it(
    cleanup_dtos, monkeypatch
):
    """The limit check must use the already-known DTO.size, not len() on
    file content -- read_bounded already enforced this during the upload
    read, so this is a defensive check against a future caller building a
    FileUploadDTO some other way, and it must not require materialising
    (or even opening) the file to run.
    """
    dto = _dto(b"tiny actual content", filename="huge.txt", content_type="text/plain")
    dto = FileUploadDTO(
        filename=dto.filename,
        content_type=dto.content_type,
        size=MAX_FILE_SIZE_BYTES + 1,
        path=dto.path,
        content_hash=dto.content_hash,
    )
    cleanup_dtos.append(dto)

    def _fail_if_opened(*_args: object, **_kwargs: object) -> None:
        raise AssertionError("file_too_large check must not open the file")

    monkeypatch.setattr(Path, "open", _fail_if_opened)

    with pytest.raises(ApplicationError) as exc_info:
        await process_document(dto)

    assert exc_info.value.status_code == 413
    assert exc_info.value.code == "document.file_too_large"


@pytest.mark.asyncio
async def test_builds_the_blob_from_path_rather_than_reading_data_up_front(
    cleanup_dtos, monkeypatch
):
    dto = _dto(b"hello", filename="a.txt", content_type="text/plain")
    cleanup_dtos.append(dto)

    captured: dict[str, object] = {}
    real_from_path = Blob.from_path

    def _spy_from_path(path: Path, **kwargs: object):
        blob = real_from_path(path, **kwargs)
        captured["blob"] = blob
        return blob

    monkeypatch.setattr(
        "src.documents.processor.Blob.from_path", _spy_from_path
    )

    await process_document(dto)

    assert "blob" in captured
    assert captured["blob"].data is None
    assert captured["blob"].path == dto.path


@pytest.mark.asyncio
async def test_docx_parser_reads_directly_from_path_not_full_bytes(
    cleanup_dtos, monkeypatch
):
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_paragraph("A real docx paragraph.")
    buf = io.BytesIO()
    doc.save(buf)

    dto = _dto(
        buf.getvalue(),
        filename="a.docx",
        content_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
    )
    cleanup_dtos.append(dto)

    def _fail_as_bytes(self):
        raise AssertionError(
            "docx parser must read from blob.path, not blob.as_bytes()"
        )

    monkeypatch.setattr(Blob, "as_bytes", _fail_as_bytes)

    docs = await process_document(dto)

    assert any("A real docx paragraph." in d.page_content for d in docs)


@pytest.mark.asyncio
async def test_xlsx_parser_reads_directly_from_path_not_full_bytes(
    cleanup_dtos, monkeypatch
):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["name", "age"])
    ws.append(["Ada", 36])
    buf = io.BytesIO()
    wb.save(buf)

    dto = _dto(
        buf.getvalue(),
        filename="a.xlsx",
        content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )
    cleanup_dtos.append(dto)

    def _fail_as_bytes(self):
        raise AssertionError(
            "xlsx parser must read from blob.path, not blob.as_bytes()"
        )

    monkeypatch.setattr(Blob, "as_bytes", _fail_as_bytes)

    docs = await process_document(dto)

    assert any("Ada" in d.page_content for d in docs)


@pytest.mark.asyncio
async def test_pptx_parser_reads_directly_from_path_not_full_bytes(
    cleanup_dtos, monkeypatch
):
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "A real slide title"
    buf = io.BytesIO()
    prs.save(buf)

    dto = _dto(
        buf.getvalue(),
        filename="a.pptx",
        content_type=(
            "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        ),
    )
    cleanup_dtos.append(dto)

    def _fail_as_bytes(self):
        raise AssertionError(
            "pptx parser must read from blob.path, not blob.as_bytes()"
        )

    monkeypatch.setattr(Blob, "as_bytes", _fail_as_bytes)

    docs = await process_document(dto)

    assert any("A real slide title" in d.page_content for d in docs)
