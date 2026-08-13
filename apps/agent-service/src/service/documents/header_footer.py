"""Header/footer content rendering: `{page}`/`{pages}`/`{title}`/`{date}`/
`{version}` placeholders, laid out in a borderless 3-column table.

A table — not tab stops — positions the left/center/right segments: custom
`w:tabs` stop positions are not reliably honored by every DOCX viewer (in
particular the browser-based renderer this project's own web frontend uses
to preview generated files), while table columns render consistently
everywhere. `{page}`/`{pages}` become real Word fields (via
`ooxml.add_field_run`) since python-docx cannot know the final page count;
the other placeholders are plain text substitution from values the caller
already has.
"""

from __future__ import annotations

import re
from typing import TYPE_CHECKING, Any

from service.documents.ooxml import add_field_run, add_table_border

if TYPE_CHECKING:
    from service.documents.options import HeaderFooterOptions

_PLACEHOLDER_RE = re.compile(r"\{(page|pages|title|date|version)\}")
_FIELD_CODES = {"page": "PAGE", "pages": "NUMPAGES"}

# Center gets more room than the flanking columns — the common case is a
# short label left/right and a longer "Sayfa {page} / {pages}" in the middle.
_COLUMN_WIDTH_RATIOS = (0.25, 0.5, 0.25)


def has_content(options: HeaderFooterOptions) -> bool:
    return bool(options.left or options.center or options.right or options.rule)


def _render_segment(paragraph: Any, template: str, substitutions: dict[str, str]) -> None:
    pos = 0
    for match in _PLACEHOLDER_RE.finditer(template):
        if match.start() > pos:
            paragraph.add_run(template[pos : match.start()])
        token = match.group(1)
        if token in _FIELD_CODES:
            add_field_run(paragraph, _FIELD_CODES[token])
        else:
            paragraph.add_run(substitutions.get(token, ""))
        pos = match.end()
    if pos < len(template):
        paragraph.add_run(template[pos:])


def render(
    container: Any,
    options: HeaderFooterOptions,
    substitutions: dict[str, str],
    content_width_emu: int,
    border_side: str,
) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Emu

    table = container.add_table(rows=1, cols=3, width=Emu(content_width_emu))
    table.autofit = False
    alignments = (WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT)
    segments = (options.left, options.center, options.right)

    for cell, ratio, alignment, template in zip(
        table.rows[0].cells, _COLUMN_WIDTH_RATIOS, alignments, segments, strict=True
    ):
        cell.width = Emu(int(content_width_emu * ratio))
        paragraph = cell.paragraphs[0]
        paragraph.alignment = alignment
        _render_segment(paragraph, template, substitutions)

    if options.rule:
        add_table_border(table, side=border_side)
