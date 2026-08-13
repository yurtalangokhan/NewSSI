"""DOCX rendering from the parsed block tree.

Every block becomes Word paragraphs/tables/inline shapes with real inline
run formatting (Group C), plus the full page scaffolding and theme
application (Group A/B): page size/orientation/margins, header/footer with
{page}/{pages} fields, cover page, table of contents, heading numbering,
and theme fonts/colors on headings, body text, links, table headers/zebra
rows and code block backgrounds.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from typing import TYPE_CHECKING, Any

from service.documents import header_footer
from service.documents.blocks import (
    Block,
    BulletListBlock,
    CodeBlock,
    FootnoteDefBlock,
    HeadingBlock,
    HorizontalRuleBlock,
    ImageBlock,
    InlineSpan,
    ListItem,
    OrderedListBlock,
    PageBreakBlock,
    ParagraphBlock,
    QuoteBlock,
    TableBlock,
)
from service.documents.image_loading import load_image_bytes
from service.documents.numbering import prepare_blocks
from service.documents.ooxml import (
    add_bookmark,
    add_field_begin,
    add_field_end,
    add_horizontal_rule,
    add_hyperlink_run,
    add_internal_hyperlink_run,
    enable_update_fields,
    shade_cell,
    shade_paragraph,
)
from service.documents.options import CoverOptions, DocumentOptions, TableStyleOptions, TocOptions
from service.documents.themes import EffectiveStyle, resolve_effective_style

if TYPE_CHECKING:
    from docx.text.paragraph import Paragraph

_MONO_FONT = "Consolas"
_IMAGE_WIDTH_INCHES = 6.0

_BULLET_STYLES = ("List Bullet", "List Bullet 2", "List Bullet 3")
_ORDERED_STYLES = ("List Number", "List Number 2", "List Number 3")

_PAGE_SIZES_MM = {"A4": (210.0, 297.0), "Letter": (215.9, 279.4)}


@dataclass(frozen=True)
class _RunStyle:
    font: str
    color: str
    link_color: str


def render_docx(title: str | None, markdown: str, options: DocumentOptions | None = None) -> bytes:
    import docx

    if not (markdown or "").strip():
        raise ValueError("render_docx: content must not be empty")

    options = options or DocumentOptions()
    style = resolve_effective_style(options)
    blocks = prepare_blocks(markdown, options)
    heading_bookmark_names = _heading_bookmark_names(blocks)

    document = docx.Document()
    _apply_page_setup(document, options.page.size, options.page.orientation, options.page.margins)
    _apply_header_footer(document, options, title)

    heading_style = _RunStyle(style.heading_font, style.heading_color, style.link_color)
    body_style = _RunStyle(style.body_font, style.text_color, style.link_color)

    if options.cover.enabled:
        _add_cover_page(document, options.cover, title, heading_style)
    elif title:
        paragraph = document.add_paragraph(title, style="Title")
        _apply_run_style(paragraph.runs, heading_style)

    if options.toc.enabled:
        _add_toc(document, options.toc, blocks, body_style, heading_bookmark_names)

    heading_index = 0
    for block in blocks:
        bookmark = None
        if isinstance(block, HeadingBlock):
            bookmark = (heading_bookmark_names[heading_index], heading_index)
            heading_index += 1
        _add_docx_block(document, block, style, heading_style, body_style, options.tables, bookmark)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _apply_page_setup(document: Any, size: str, orientation: str, margins: Any) -> None:
    from docx.enum.section import WD_ORIENT
    from docx.shared import Mm

    section = document.sections[0]
    width_mm, height_mm = _PAGE_SIZES_MM[size]
    if orientation == "landscape":
        width_mm, height_mm = height_mm, width_mm
        section.orientation = WD_ORIENT.LANDSCAPE
    else:
        section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Mm(width_mm)
    section.page_height = Mm(height_mm)
    section.top_margin = Mm(margins.top)
    section.bottom_margin = Mm(margins.bottom)
    section.left_margin = Mm(margins.left)
    section.right_margin = Mm(margins.right)


def _apply_header_footer(document: Any, options: DocumentOptions, title: str | None) -> None:
    section = document.sections[0]
    content_width = section.page_width - section.left_margin - section.right_margin
    substitutions = {
        "title": title or options.cover.project or "",
        "date": options.cover.date,
        "version": options.cover.version,
    }

    header_active = header_footer.has_content(options.header)
    footer_active = header_footer.has_content(options.footer)
    if not header_active and not footer_active:
        return

    if (header_active and options.header.different_first_page) or (
        footer_active and options.footer.different_first_page
    ):
        section.different_first_page_header_footer = True

    if header_active:
        header_footer.render(section.header, options.header, substitutions, content_width, "bottom")
    if footer_active:
        header_footer.render(section.footer, options.footer, substitutions, content_width, "top")


_COVER_LOGO_WIDTH_INCHES = 1.6


def _add_cover_page(
    document: Any, cover: CoverOptions, title: str | None, heading_style: _RunStyle
) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    def centered_paragraph(
        text: str, *, bold: bool = False, italic: bool = False, size: int | None = None
    ):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.bold = bold
        run.italic = italic
        if size:
            run.font.size = Pt(size)
        return paragraph

    if cover.logo:
        data = load_image_bytes(cover.logo)
        if data is not None:
            picture_paragraph = document.add_paragraph()
            picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture_paragraph.add_run().add_picture(
                io.BytesIO(data), width=Inches(_COVER_LOGO_WIDTH_INCHES)
            )

    if title:
        title_paragraph = document.add_paragraph(title, style="Title")
        title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _apply_run_style(title_paragraph.runs, heading_style)
    if cover.subtitle:
        centered_paragraph(cover.subtitle, italic=True, size=16)

    document.add_paragraph()
    document.add_paragraph()

    for value in (
        cover.project,
        f"v{cover.version}" if cover.version else "",
        cover.date,
        cover.author,
        cover.organization,
    ):
        if value:
            centered_paragraph(value)

    if cover.classification:
        centered_paragraph(cover.classification, bold=True)

    document.add_page_break()


_TOC_ENTRY_INDENT_INCHES = 0.25


def _heading_bookmark_names(blocks: list[Block]) -> list[str]:
    """One bookmark name per heading block, in document order — shared by
    the TOC (which links to these) and the heading paragraphs themselves
    (which anchor them), so a TOC entry always points at the right heading
    regardless of `toc.depth` filtering."""
    count = sum(1 for block in blocks if isinstance(block, HeadingBlock))
    return [f"_Toc_h{i}" for i in range(count)]


def _add_toc(
    document: Any,
    toc: TocOptions,
    blocks: list[Block],
    body_style: _RunStyle,
    heading_bookmark_names: list[str],
) -> None:
    """A real Word TOC field (so Word computes real page numbers on open)
    with a cached, unnumbered-page preview in between — see design note in
    the enrichment spec §6.2 on why viewers that don't evaluate fields (e.g.
    docx-preview) need this cached content to show anything at all. Each
    cached entry is also a real hyperlink to its heading's bookmark, so
    navigation works even in those viewers.
    """
    from docx.shared import Inches

    entries: list[tuple[int, str, str]] = []
    heading_index = 0
    for block in blocks:
        if not isinstance(block, HeadingBlock):
            continue
        bookmark_name = heading_bookmark_names[heading_index]
        heading_index += 1
        if block.level <= toc.depth:
            entries.append((block.level, block.text, bookmark_name))
    if not entries:
        return

    enable_update_fields(document)

    title_paragraph = document.add_paragraph(toc.title, style="TOC Heading")
    _apply_run_style(title_paragraph.runs, body_style)

    field_paragraph = document.add_paragraph()
    add_field_begin(field_paragraph, f'TOC \\o "1-{toc.depth}" \\h \\z \\u')

    for level, text, bookmark_name in entries:
        entry_paragraph = document.add_paragraph()
        entry_paragraph.paragraph_format.left_indent = Inches(
            _TOC_ENTRY_INDENT_INCHES * (level - 1)
        )
        add_internal_hyperlink_run(entry_paragraph, text, anchor=bookmark_name)

    add_field_end(document.add_paragraph())
    document.add_page_break()


def _apply_run_style(runs: Any, run_style: _RunStyle) -> None:
    from docx.shared import RGBColor

    for run in runs:
        run.font.name = run_style.font
        run.font.color.rgb = RGBColor.from_string(run_style.color.lstrip("#"))


def _add_span_run(paragraph: Paragraph, span: InlineSpan, run_style: _RunStyle) -> None:
    from docx.shared import RGBColor

    if span.footnote_id:
        run = paragraph.add_run(f"[{span.footnote_id}]")
        run.font.superscript = True
        return
    if span.link:
        add_hyperlink_run(paragraph, span.text, span.link, color=run_style.link_color.lstrip("#"))
        return
    run = paragraph.add_run(span.text)
    if span.bold:
        run.bold = True
    if span.italic:
        run.italic = True
    if span.code:
        run.font.name = _MONO_FONT
    else:
        run.font.name = run_style.font
        run.font.color.rgb = RGBColor.from_string(run_style.color.lstrip("#"))


def _add_spans(paragraph: Paragraph, spans: list[InlineSpan], run_style: _RunStyle) -> None:
    for span in spans:
        _add_span_run(paragraph, span, run_style)


def _list_style_name(is_ordered: bool, level: int) -> str:
    styles = _ORDERED_STYLES if is_ordered else _BULLET_STYLES
    return styles[min(level, len(styles) - 1)]


def _add_list_item(document: Any, item: ListItem, is_ordered: bool, run_style: _RunStyle) -> None:
    paragraph = document.add_paragraph(style=_list_style_name(is_ordered, item.level))
    if item.checked is not None:
        paragraph.add_run("☑ " if item.checked else "☐ ")
    _add_spans(paragraph, item.spans, run_style)


def _add_table(
    document: Any, block: TableBlock, style: EffectiveStyle, table_options: TableStyleOptions
) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    alignment_map = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }
    body_run_style = _RunStyle(style.body_font, style.text_color, style.link_color)
    header_run_style = _RunStyle(style.body_font, style.table_header_text, style.link_color)

    table = document.add_table(rows=0, cols=len(block.header))
    table.style = "Table Grid" if table_options.style != "minimal" else "Table Normal"

    for row_index, row_cells_source in enumerate((block.header, *block.rows)):
        is_header = row_index == 0
        row_cells = table.add_row().cells
        for index, (cell, source) in enumerate(zip(row_cells, row_cells_source, strict=False)):
            paragraph = cell.paragraphs[0]
            _add_spans(paragraph, source.spans, header_run_style if is_header else body_run_style)
            if is_header:
                if table_options.header_bold:
                    for run in paragraph.runs:
                        run.bold = True
                shade_cell(cell, style.table_header_bg)
            elif table_options.style == "zebra" and row_index % 2 == 0:
                shade_cell(cell, style.table_zebra_bg)
            if block.alignments and index < len(block.alignments):
                paragraph.alignment = alignment_map.get(block.alignments[index])

    if block.caption:
        caption_paragraph = document.add_paragraph(style="Caption")
        caption_paragraph.add_run(block.caption)


def _add_image(document: Any, block: ImageBlock) -> None:
    from docx.shared import Inches

    data = load_image_bytes(block.src)
    if data is None:
        return
    document.add_picture(io.BytesIO(data), width=Inches(_IMAGE_WIDTH_INCHES))
    if block.caption:
        caption_paragraph = document.add_paragraph(style="Caption")
        caption_paragraph.add_run(block.caption)


def _add_quote(document: Any, block: QuoteBlock, run_style: _RunStyle) -> None:
    from docx.shared import Inches

    paragraph = document.add_paragraph(style="Quote")
    paragraph.paragraph_format.left_indent = Inches(0.4)
    if block.callout_type:
        paragraph.add_run(f"{block.callout_type}: ").bold = True
    _add_spans(paragraph, block.spans, run_style)


def _add_footnote_def(document: Any, block: FootnoteDefBlock, run_style: _RunStyle) -> None:
    paragraph = document.add_paragraph()
    paragraph.add_run(f"[{block.id}] ").bold = True
    _add_spans(paragraph, block.spans, run_style)


def _add_docx_block(
    document: Any,
    block: Block,
    style: EffectiveStyle,
    heading_style: _RunStyle,
    body_style: _RunStyle,
    table_options: TableStyleOptions,
    bookmark: tuple[str, int] | None = None,
) -> None:
    if isinstance(block, HeadingBlock):
        paragraph = document.add_paragraph(style=f"Heading {block.level}")
        if bookmark:
            bookmark_name, bookmark_id = bookmark
            add_bookmark(paragraph, bookmark_name, bookmark_id)
        _add_spans(paragraph, block.spans, heading_style)
    elif isinstance(block, ParagraphBlock):
        _add_spans(document.add_paragraph(), block.spans, body_style)
    elif isinstance(block, BulletListBlock):
        for item in block.items:
            _add_list_item(document, item, is_ordered=False, run_style=body_style)
    elif isinstance(block, OrderedListBlock):
        for item in block.items:
            _add_list_item(document, item, is_ordered=True, run_style=body_style)
    elif isinstance(block, CodeBlock):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(block.text)
        run.font.name = _MONO_FONT
        shade_paragraph(paragraph, style.code_bg)
    elif isinstance(block, TableBlock):
        _add_table(document, block, style, table_options)
    elif isinstance(block, QuoteBlock):
        _add_quote(document, block, body_style)
    elif isinstance(block, ImageBlock):
        _add_image(document, block)
    elif isinstance(block, HorizontalRuleBlock):
        add_horizontal_rule(document.add_paragraph())
    elif isinstance(block, PageBreakBlock):
        document.add_page_break()
    elif isinstance(block, FootnoteDefBlock):
        _add_footnote_def(document, block, body_style)
