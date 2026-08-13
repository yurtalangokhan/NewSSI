"""XLSX / CSV rendering.

Styling (Group E) applies uniformly to every sheet in the workbook from one
`SpreadsheetOptions`: a themed header row, freeze pane, autofilter, zebra
striping, column widths/number-formats/alignments and static conditional
cell shading. "Static" because these are computed once at generation time
from the values already known — not a live Excel conditional-formatting
rule that would recalculate if the user edits the sheet afterwards.
"""

from __future__ import annotations

import csv
import io
import re
from typing import Any

from service.documents.options import ConditionalFormatRule, SpreadsheetOptions
from service.documents.themes import ThemeDefinition, resolve_theme

_INVALID_SHEET_NAME_CHARS_RE = re.compile(r"[\\/*?:\[\]]")
_MAX_SHEET_NAME_LENGTH = 31

_NUMBER_FORMATS = {
    "text": "@",
    "number": "0",
    "date": "yyyy-mm-dd",
    "percent": "0%",
    "currency": "#,##0.00",
}
_MIN_AUTOFIT_WIDTH = 8
_MAX_AUTOFIT_WIDTH = 60
_AUTOFIT_PADDING = 2


def _sanitize_sheet_name(name: str) -> str:
    sanitized = _INVALID_SHEET_NAME_CHARS_RE.sub("_", name)
    return sanitized[:_MAX_SHEET_NAME_LENGTH] or "Sheet"


def render_xlsx(sheets: list[dict[str, Any]], options: SpreadsheetOptions | None = None) -> bytes:
    import openpyxl

    if not sheets:
        raise ValueError("render_xlsx requires at least one sheet")

    options = options or SpreadsheetOptions()
    theme = resolve_theme(options.theme)

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    for spec in sheets:
        rows = spec.get("rows") or []
        if not rows:
            raise ValueError(f"Sheet '{spec.get('name', '')}' has no rows")
        worksheet = workbook.create_sheet(title=_sanitize_sheet_name(spec.get("name") or "Sheet"))
        for row in rows:
            worksheet.append(list(row))
        _style_sheet(worksheet, rows, options, theme)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _style_sheet(
    worksheet: Any, rows: list[list[Any]], options: SpreadsheetOptions, theme: ThemeDefinition
) -> None:
    n_rows = len(rows)
    n_cols = max((len(row) for row in rows), default=0)
    first_data_row = 2 if options.header_row else 1

    if options.header_row and n_rows:
        _style_header_row(worksheet, n_cols, theme)
    if options.freeze_header and options.header_row and n_rows:
        worksheet.freeze_panes = "A2"
    if options.autofilter and n_rows and n_cols:
        from openpyxl.utils import get_column_letter

        worksheet.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{n_rows}"
    if options.zebra:
        _apply_zebra_striping(worksheet, n_rows, n_cols, first_data_row, theme)
    if options.autofit_columns:
        _autofit_columns(worksheet, rows, n_cols)
    if options.column_widths:
        _apply_column_widths(worksheet, options.column_widths)
    if options.column_formats:
        _apply_column_formats(worksheet, options.column_formats, n_rows, first_data_row)
    if options.column_alignments:
        _apply_column_alignments(worksheet, options.column_alignments, n_rows, first_data_row)
    for rule in options.conditional_formats:
        _apply_conditional_format(worksheet, rule, n_rows, first_data_row)


def _style_header_row(worksheet: Any, n_cols: int, theme: ThemeDefinition) -> None:
    from openpyxl.styles import Font, PatternFill

    fill = PatternFill(
        start_color=theme.table_header_bg.lstrip("#"),
        end_color=theme.table_header_bg.lstrip("#"),
        fill_type="solid",
    )
    font = Font(bold=True, color=theme.table_header_text.lstrip("#"))
    for col in range(1, n_cols + 1):
        cell = worksheet.cell(row=1, column=col)
        cell.fill = fill
        cell.font = font


def _apply_zebra_striping(
    worksheet: Any, n_rows: int, n_cols: int, first_data_row: int, theme: ThemeDefinition
) -> None:
    from openpyxl.styles import PatternFill

    fill = PatternFill(
        start_color=theme.table_zebra_bg.lstrip("#"),
        end_color=theme.table_zebra_bg.lstrip("#"),
        fill_type="solid",
    )
    for offset, row in enumerate(range(first_data_row, n_rows + 1)):
        if offset % 2 == 1:
            for col in range(1, n_cols + 1):
                worksheet.cell(row=row, column=col).fill = fill


def _autofit_columns(worksheet: Any, rows: list[list[Any]], n_cols: int) -> None:
    from openpyxl.utils import get_column_letter

    for col in range(1, n_cols + 1):
        longest = max((len(str(row[col - 1])) for row in rows if col - 1 < len(row)), default=0)
        width = max(_MIN_AUTOFIT_WIDTH, min(longest + _AUTOFIT_PADDING, _MAX_AUTOFIT_WIDTH))
        worksheet.column_dimensions[get_column_letter(col)].width = width


def _apply_column_widths(worksheet: Any, widths: list[int]) -> None:
    from openpyxl.utils import get_column_letter

    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width


def _apply_column_formats(
    worksheet: Any, formats: list[str], n_rows: int, first_data_row: int
) -> None:
    for index, format_name in enumerate(formats, start=1):
        excel_format = _NUMBER_FORMATS.get(format_name)
        if not excel_format:
            continue
        for row in range(first_data_row, n_rows + 1):
            worksheet.cell(row=row, column=index).number_format = excel_format


def _apply_column_alignments(
    worksheet: Any, alignments: list[str], n_rows: int, first_data_row: int
) -> None:
    from openpyxl.styles import Alignment

    for index, alignment_name in enumerate(alignments, start=1):
        for row in range(first_data_row, n_rows + 1):
            worksheet.cell(row=row, column=index).alignment = Alignment(horizontal=alignment_name)


def _matches_conditional_rule(rule: ConditionalFormatRule, value: Any) -> bool:
    if rule.rule == "equals":
        return value == rule.value
    if rule.rule == "greater_than":
        return (
            isinstance(value, int | float)
            and isinstance(rule.value, int | float)
            and value > rule.value
        )
    if rule.rule == "less_than":
        return (
            isinstance(value, int | float)
            and isinstance(rule.value, int | float)
            and value < rule.value
        )
    if rule.rule == "contains":
        return isinstance(value, str) and isinstance(rule.value, str) and rule.value in value
    return False


def _apply_conditional_format(
    worksheet: Any, rule: ConditionalFormatRule, n_rows: int, first_data_row: int
) -> None:
    from openpyxl.styles import PatternFill

    fill = PatternFill(
        start_color=rule.color.lstrip("#"), end_color=rule.color.lstrip("#"), fill_type="solid"
    )
    for row in range(first_data_row, n_rows + 1):
        cell = worksheet.cell(row=row, column=rule.column)
        if _matches_conditional_rule(rule, cell.value):
            cell.fill = fill


def render_csv(rows: list[list[Any]], options: SpreadsheetOptions | None = None) -> bytes:
    options = options or SpreadsheetOptions()
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=options.delimiter)
    for row in rows:
        writer.writerow(["" if cell is None else cell for cell in row])
    encoded = buffer.getvalue().encode("utf-8")
    # UTF-8 BOM so Excel opens Turkish characters correctly; off by request
    # for consumers (e.g. some *nix tooling) that choke on a leading BOM.
    return (b"\xef\xbb\xbf" + encoded) if options.csv_bom else encoded
