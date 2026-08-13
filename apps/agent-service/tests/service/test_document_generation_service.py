"""Tests for DocumentGenerationService — markdown parsing and format renderers."""

import csv
import io
import zipfile

import pytest

from service.DocumentGenerationService import (
    BulletListBlock,
    CodeBlock,
    HeadingBlock,
    OrderedListBlock,
    ParagraphBlock,
    TableBlock,
    mime_for_format,
    parse_markdown_blocks,
    render_csv,
    render_docx,
    render_markdown_text,
    render_pdf,
    render_xlsx,
    sanitize_filename,
)

TURKISH = "Çalışma özeti: ğüşıöç İĞÜŞİÖÇ"


# ---------------------------------------------------------------------------
# Markdown parsing
# ---------------------------------------------------------------------------


def test_parses_headings_with_level():
    blocks = parse_markdown_blocks("# Başlık\n\n### Alt başlık")

    assert [(b.level, b.text) for b in blocks] == [(1, "Başlık"), (3, "Alt başlık")]
    assert all(isinstance(b, HeadingBlock) for b in blocks)


def test_parses_consecutive_lines_into_one_paragraph():
    blocks = parse_markdown_blocks("ilk satır\nikinci satır\n\nayrı paragraf")

    assert [b.text for b in blocks] == ["ilk satır ikinci satır", "ayrı paragraf"]
    assert all(isinstance(b, ParagraphBlock) for b in blocks)


def test_parses_bullet_list():
    [block] = parse_markdown_blocks("- elma\n* armut\n- erik")

    assert isinstance(block, BulletListBlock)
    assert [i.text for i in block.items] == ["elma", "armut", "erik"]


def test_parses_ordered_list():
    [block] = parse_markdown_blocks("1. birinci\n2. ikinci")

    assert isinstance(block, OrderedListBlock)
    assert [i.text for i in block.items] == ["birinci", "ikinci"]


def test_parses_pipe_table_with_header_and_rows():
    markdown = "| Ürün | Adet |\n| --- | --- |\n| Elma | 3 |\n| Armut | 5 |"

    [block] = parse_markdown_blocks(markdown)

    assert isinstance(block, TableBlock)
    assert [c.text for c in block.header] == ["Ürün", "Adet"]
    assert [[c.text for c in row] for row in block.rows] == [["Elma", "3"], ["Armut", "5"]]


def test_parses_fenced_code_block_verbatim():
    markdown = "```python\nx = 1\n\ny = 2\n```"

    blocks = parse_markdown_blocks(markdown)

    assert blocks == [CodeBlock(text="x = 1\n\ny = 2", language="python")]


def test_preserves_inline_emphasis_as_spans_but_text_stays_plain():
    """Inline markers are no longer stripped (Group C) — they become spans
    that renderers turn into real formatting. `.text` still gives plain text."""
    [block] = parse_markdown_blocks("bu **kalın** ve *eğik* ve `kod`")

    assert block.text == "bu kalın ve eğik ve kod"
    assert any(s.bold for s in block.spans)
    assert any(s.italic for s in block.spans)
    assert any(s.code for s in block.spans)


def test_returns_no_blocks_for_blank_markdown():
    assert parse_markdown_blocks("   \n\n  ") == []


# ---------------------------------------------------------------------------
# Filename / mime helpers
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("raw", "fmt", "expected"),
    [
        ("rapor", "pdf", "rapor.pdf"),
        ("rapor.pdf", "pdf", "rapor.pdf"),
        ("rapor.PDF", "pdf", "rapor.PDF"),
        ("../../etc/passwd", "csv", "etc_passwd.csv"),
        ("a/b\\c", "docx", "a_b_c.docx"),
        ("", "xlsx", "document.xlsx"),
        ("   ", "xlsx", "document.xlsx"),
    ],
)
def test_sanitize_filename(raw, fmt, expected):
    assert sanitize_filename(raw, fmt) == expected


def test_sanitize_filename_truncates_long_names_but_keeps_extension():
    result = sanitize_filename("a" * 300, "pdf")

    assert len(result) <= 120
    assert result.endswith(".pdf")


@pytest.mark.parametrize(
    ("fmt", "expected"),
    [
        ("pdf", "application/pdf"),
        ("docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        ("csv", "text/csv"),
        ("md", "text/markdown"),
        ("txt", "text/plain"),
    ],
)
def test_mime_for_format(fmt, expected):
    assert mime_for_format(fmt) == expected


def test_mime_for_unknown_format_raises():
    with pytest.raises(ValueError, match="rtf"):
        mime_for_format("rtf")


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------


def test_render_csv_round_trips_rows_including_turkish():
    data = render_csv([["Ürün", "Adet"], ["Çilek", 3]])

    text = data.decode("utf-8-sig")
    assert list(csv.reader(io.StringIO(text))) == [["Ürün", "Adet"], ["Çilek", "3"]]


def test_render_csv_writes_utf8_bom_for_excel_compatibility():
    assert render_csv([["a"]]).startswith(b"\xef\xbb\xbf")


def test_render_csv_renders_none_as_empty_cell():
    text = render_csv([["a", None]]).decode("utf-8-sig")

    assert list(csv.reader(io.StringIO(text))) == [["a", ""]]


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------


def test_render_xlsx_creates_named_sheets_with_values():
    import openpyxl

    data = render_xlsx(
        [
            {"name": "Satışlar", "rows": [["Ürün", "Adet"], ["Çilek", 3]]},
            {"name": "Notlar", "rows": [["ğüşıöç"]]},
        ]
    )

    workbook = openpyxl.load_workbook(io.BytesIO(data))
    assert workbook.sheetnames == ["Satışlar", "Notlar"]
    assert list(workbook["Satışlar"].iter_rows(values_only=True)) == [
        ("Ürün", "Adet"),
        ("Çilek", 3),
    ]
    assert list(workbook["Notlar"].iter_rows(values_only=True)) == [("ğüşıöç",)]


def test_render_xlsx_sanitizes_invalid_sheet_names():
    import openpyxl

    data = render_xlsx([{"name": "a/b:c*d?e[f]", "rows": [["x"]]}])

    workbook = openpyxl.load_workbook(io.BytesIO(data))
    assert workbook.sheetnames == ["a_b_c_d_e_f_"]


def test_render_xlsx_requires_at_least_one_sheet():
    with pytest.raises(ValueError, match="at least one sheet"):
        render_xlsx([])


def test_render_xlsx_rejects_sheet_without_rows():
    with pytest.raises(ValueError, match="rows"):
        render_xlsx([{"name": "Bos", "rows": []}])


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def test_render_docx_produces_openable_document_with_turkish_text():
    import docx

    data = render_docx("Rapor", f"# Başlık\n\n{TURKISH}")

    document = docx.Document(io.BytesIO(data))
    texts = [p.text for p in document.paragraphs]
    assert "Rapor" in texts
    assert "Başlık" in texts
    assert TURKISH in texts


def test_render_docx_renders_markdown_table_as_word_table():
    import docx

    data = render_docx(None, "| A | B |\n| --- | --- |\n| 1 | 2 |")

    document = docx.Document(io.BytesIO(data))
    assert len(document.tables) == 1
    table = document.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["A", "B"]
    assert [c.text for c in table.rows[1].cells] == ["1", "2"]


def test_render_docx_is_a_zip_container():
    assert render_docx(None, "merhaba").startswith(b"PK")


def test_render_docx_rejects_empty_content():
    with pytest.raises(ValueError, match="empty"):
        render_docx(None, "   ")


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def test_render_pdf_produces_pdf_magic_bytes():
    assert render_pdf("Rapor", "merhaba").startswith(b"%PDF")


def test_render_pdf_embeds_turkish_text_extractably():
    from pypdf import PdfReader

    data = render_pdf("Rapor", f"# Başlık\n\n{TURKISH}")

    text = PdfReader(io.BytesIO(data)).pages[0].extract_text()
    assert "Başlık" in text
    assert "ğüşıöç" in text


def test_render_pdf_renders_tables_and_lists_without_error():
    from pypdf import PdfReader

    markdown = "| A | B |\n| --- | --- |\n| 1 | 2 |\n\n- madde bir\n- madde iki"

    text = PdfReader(io.BytesIO(render_pdf(None, markdown))).pages[0].extract_text()

    assert "madde bir" in text
    assert "A" in text


def test_render_pdf_rejects_empty_content():
    with pytest.raises(ValueError, match="empty"):
        render_pdf(None, "")


# ---------------------------------------------------------------------------
# Markdown / plain text
# ---------------------------------------------------------------------------


def test_render_markdown_text_prepends_title_as_heading():
    data = render_markdown_text("Rapor", "gövde")

    assert data.decode("utf-8") == "# Rapor\n\ngövde"


def test_render_markdown_text_without_title_returns_content_unchanged():
    assert render_markdown_text(None, "gövde").decode("utf-8") == "gövde"


def test_render_markdown_text_rejects_empty_content():
    with pytest.raises(ValueError, match="empty"):
        render_markdown_text("Rapor", "")


# ---------------------------------------------------------------------------
# Cross-format smoke check
# ---------------------------------------------------------------------------


def test_docx_and_xlsx_are_valid_zip_archives():
    assert zipfile.is_zipfile(io.BytesIO(render_docx(None, "x")))
    assert zipfile.is_zipfile(io.BytesIO(render_xlsx([{"name": "S", "rows": [["x"]]}])))
