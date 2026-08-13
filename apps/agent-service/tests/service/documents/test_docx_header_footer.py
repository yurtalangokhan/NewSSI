"""Tests for DOCX header/footer rendering: a borderless 3-column table for
left/center/right positioning (tab stops are not reliably honored by every
DOCX viewer — see the design note in docx_renderer._apply_header_footer),
{page}/{pages}/{title}/{date}/{version} placeholders, the rule border and
the different-first-page (blank cover) behavior."""

import io

import docx
from docx.oxml.ns import qn

from service.documents.docx_renderer import render_docx
from service.documents.options import parse_document_options


def _cells(container):
    return container.tables[0].rows[0].cells


def test_no_header_or_footer_options_leaves_default_blank():
    data = render_docx(None, "x")

    section = docx.Document(io.BytesIO(data)).sections[0]

    assert section.header.tables == []
    assert section.footer.tables == []


def test_header_left_and_right_text_appear_in_their_columns():
    options = parse_document_options({"header": {"left": "AgenticAI", "right": "SRS-001"}})

    data = render_docx(None, "x", options=options)

    cells = _cells(docx.Document(io.BytesIO(data)).sections[0].header)
    assert cells[0].text == "AgenticAI"
    assert cells[2].text == "SRS-001"


def test_footer_page_and_pages_render_as_real_word_fields():
    options = parse_document_options({"footer": {"center": "Sayfa {page} / {pages}"}})

    data = render_docx(None, "x", options=options)

    center_cell = _cells(docx.Document(io.BytesIO(data)).sections[0].footer)[1]
    instr_texts = [el.text for el in center_cell._tc.findall(f".//{qn('w:instrText')}")]

    assert any("PAGE" in (t or "") for t in instr_texts)
    assert any("NUMPAGES" in (t or "") for t in instr_texts)
    assert "Sayfa" in center_cell.text
    assert "/" in center_cell.text


def test_header_rule_adds_table_bottom_border():
    options = parse_document_options({"header": {"left": "x", "rule": True}})

    data = render_docx(None, "x", options=options)

    table = docx.Document(io.BytesIO(data)).sections[0].header.tables[0]
    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    assert borders is not None
    assert borders.find(qn("w:bottom")) is not None


def test_footer_rule_adds_table_top_border():
    options = parse_document_options({"footer": {"left": "x", "rule": True}})

    data = render_docx(None, "x", options=options)

    table = docx.Document(io.BytesIO(data)).sections[0].footer.tables[0]
    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    assert borders is not None
    assert borders.find(qn("w:top")) is not None


def test_different_first_page_leaves_first_page_header_blank():
    options = parse_document_options({"header": {"left": "x", "different_first_page": True}})

    data = render_docx(None, "x", options=options)

    section = docx.Document(io.BytesIO(data)).sections[0]
    assert section.different_first_page_header_footer is True
    assert section.first_page_header.tables == []
    assert "x" in _cells(section.header)[0].text


def test_title_date_version_placeholders_are_substituted():
    options = parse_document_options(
        {
            "header": {"right": "{title} v{version} - {date}"},
            "cover": {"version": "1.0", "date": "2026-08-11"},
        }
    )

    data = render_docx("SRS", "x", options=options)

    right_cell = _cells(docx.Document(io.BytesIO(data)).sections[0].header)[2]
    assert right_cell.text == "SRS v1.0 - 2026-08-11"
