"""Tests for the DOCX table of contents: a real Word TOC field (so Word
recalculates real page numbers) with a cached heading list in between (so
`docx-preview` and other viewers that don't evaluate fields still show
something), plus the `updateFields` document setting."""

import io

import docx
from docx.oxml.ns import qn

from service.documents.docx_renderer import render_docx
from service.documents.options import parse_document_options


def _instr_texts(document) -> list[str]:
    return [el.text for el in document.element.body.findall(f".//{qn('w:instrText')}")]


def _cached_entry_texts(document) -> list[str]:
    """Plain-styled paragraph texts — the cached preview lines sit between
    the TOC Heading paragraph and the first real 'Heading N' paragraph."""
    return [
        p.text
        for p in document.paragraphs
        if p.text
        and p.style.name not in ("Heading 1", "Heading 2", "Heading 3", "TOC Heading", "Title")
    ]


def test_toc_disabled_by_default():
    data = render_docx(None, "# Giriş")

    document = docx.Document(io.BytesIO(data))
    assert not any("TOC" in (t or "") for t in _instr_texts(document))
    assert not any(p.style.name == "TOC Heading" for p in document.paragraphs)


def test_toc_field_instr_text_encodes_depth():
    options = parse_document_options({"toc": {"enabled": True, "depth": 2}})

    data = render_docx(None, "# Giriş", options=options)

    document = docx.Document(io.BytesIO(data))
    assert any('TOC \\o "1-2"' in (t or "") for t in _instr_texts(document))


def test_toc_lists_numbered_headings_before_the_real_ones():
    options = parse_document_options({"toc": {"enabled": True}, "numbering": {"headings": True}})

    data = render_docx(None, "# Giriş\n\n## Amaç", options=options)

    document = docx.Document(io.BytesIO(data))
    texts = [p.text for p in document.paragraphs]

    toc_index = texts.index("1 Giriş")
    real_heading_index = next(
        i for i, p in enumerate(document.paragraphs) if p.style.name == "Heading 1"
    )
    assert toc_index < real_heading_index
    assert "1.1 Amaç" in texts


def test_toc_respects_configured_depth():
    options = parse_document_options({"toc": {"enabled": True, "depth": 1}})

    data = render_docx(None, "# Giriş\n\n## Alt Başlık", options=options)

    document = docx.Document(io.BytesIO(data))
    cached = _cached_entry_texts(document)

    assert "Giriş" in cached
    assert "Alt Başlık" not in cached


def test_toc_title_uses_configured_text():
    options = parse_document_options({"toc": {"enabled": True, "title": "İçindekiler"}})

    data = render_docx(None, "# Giriş", options=options)

    document = docx.Document(io.BytesIO(data))
    assert any(
        p.style.name == "TOC Heading" and p.text == "İçindekiler" for p in document.paragraphs
    )


def test_toc_enables_update_fields_document_setting():
    options = parse_document_options({"toc": {"enabled": True}})

    data = render_docx(None, "# Giriş", options=options)

    document = docx.Document(io.BytesIO(data))
    update_fields = document.settings.element.find(qn("w:updateFields"))
    assert update_fields is not None
    assert update_fields.get(qn("w:val")) == "true"


def test_toc_adds_page_break_before_body():
    options = parse_document_options({"toc": {"enabled": True}})

    data = render_docx(None, "# Giriş", options=options)

    document = docx.Document(io.BytesIO(data))
    breaks = document.element.body.findall(f".//{qn('w:br')}")
    assert any(b.get(qn("w:type")) == "page" for b in breaks)


def test_toc_skipped_when_document_has_no_headings():
    options = parse_document_options({"toc": {"enabled": True}})

    data = render_docx(None, "sadece gövde metni", options=options)

    document = docx.Document(io.BytesIO(data))
    assert not any(p.style.name == "TOC Heading" for p in document.paragraphs)


# ---------------------------------------------------------------------------
# TOC entries link to their heading — works even in viewers (e.g. the web
# frontend's docx-preview) that never evaluate the TOC field itself.
# ---------------------------------------------------------------------------


def _bookmark_names(document) -> set[str]:
    return {
        el.get(qn("w:name")) for el in document.element.body.findall(f".//{qn('w:bookmarkStart')}")
    }


def test_every_heading_gets_a_bookmark():
    data = render_docx(None, "# Giriş\n\n## Amaç")

    document = docx.Document(io.BytesIO(data))
    assert len(_bookmark_names(document)) == 2


def test_toc_entry_is_a_hyperlink_anchored_to_its_heading_bookmark():
    options = parse_document_options({"toc": {"enabled": True}})

    data = render_docx(None, "# Giriş\n\n## Amaç", options=options)

    document = docx.Document(io.BytesIO(data))
    bookmark_names = _bookmark_names(document)

    toc_hyperlinks = [
        el
        for el in document.element.body.findall(f".//{qn('w:hyperlink')}")
        if el.get(qn("w:anchor"))
    ]
    assert len(toc_hyperlinks) == 2
    for hyperlink in toc_hyperlinks:
        assert hyperlink.get(qn("w:anchor")) in bookmark_names


def test_toc_hyperlink_anchor_points_to_the_matching_heading_text():
    options = parse_document_options({"toc": {"enabled": True}})

    data = render_docx(None, "# Giriş\n\n## Amaç", options=options)

    document = docx.Document(io.BytesIO(data))
    body = document.element.body

    heading_text_by_bookmark = {}
    for bookmark in body.findall(f".//{qn('w:bookmarkStart')}"):
        name = bookmark.get(qn("w:name"))
        paragraph = bookmark.getparent()
        heading_text_by_bookmark[name] = "".join(
            t.text or "" for t in paragraph.findall(f".//{qn('w:t')}")
        )

    for hyperlink in body.findall(f".//{qn('w:hyperlink')}"):
        anchor = hyperlink.get(qn("w:anchor"))
        if anchor is None:
            continue
        link_text = "".join(t.text or "" for t in hyperlink.findall(f".//{qn('w:t')}"))
        assert link_text == heading_text_by_bookmark[anchor]


def test_toc_hyperlink_anchors_only_headings_within_configured_depth():
    """The bookmark/anchor numbering must stay in sync with the TOC's own
    depth filter — a shallow toc.depth must not point entries at the wrong
    heading."""
    options = parse_document_options({"toc": {"enabled": True, "depth": 1}})

    data = render_docx(None, "# Giriş\n\n## Alt Başlık\n\n# Kapsam", options=options)

    document = docx.Document(io.BytesIO(data))
    body = document.element.body

    toc_hyperlink_texts = [
        "".join(t.text or "" for t in el.findall(f".//{qn('w:t')}"))
        for el in body.findall(f".//{qn('w:hyperlink')}")
        if el.get(qn("w:anchor"))
    ]
    assert toc_hyperlink_texts == ["Giriş", "Kapsam"]
