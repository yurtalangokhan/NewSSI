"""Tests for the DOCX renderer's Group C richness: real inline run
formatting, hyperlinks, nested/task lists, callouts, table cell formatting,
images, horizontal rules, page breaks and footnote markers."""

import io

import docx
from docx.oxml.ns import qn

from service.documents.docx_renderer import render_docx

# 1x1 transparent PNG.
_PNG_DATA_URI = (
    "data:image/png;base64,"
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
    "+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _paragraph_texts(document) -> list[str]:
    return [p.text for p in document.paragraphs]


def test_bold_italic_code_become_real_run_formatting():
    data = render_docx(None, "bu **kalın** ve *eğik* ve `kod`")

    document = docx.Document(io.BytesIO(data))
    runs = document.paragraphs[0].runs
    by_text = {r.text: r for r in runs}

    assert by_text["kalın"].bold is True
    assert by_text["eğik"].italic is True
    assert by_text["kod"].font.name == "Consolas"


def test_link_renders_as_real_hyperlink():
    data = render_docx(None, "bkz. [dokümantasyon](https://example.com/docs)")

    document = docx.Document(io.BytesIO(data))
    hyperlink_el = document.element.body.find(".//" + qn("w:hyperlink"))
    assert hyperlink_el is not None

    r_id = hyperlink_el.get(qn("r:id"))
    assert document.part.rels[r_id].target_ref == "https://example.com/docs"


def test_nested_bullet_list_uses_deeper_list_style():
    data = render_docx(None, "- üst\n  - alt")

    document = docx.Document(io.BytesIO(data))
    styles = [p.style.name for p in document.paragraphs if p.text in ("üst", "alt")]

    assert styles == ["List Bullet", "List Bullet 2"]


def test_task_list_items_show_checkbox_glyph():
    data = render_docx(None, "- [ ] yapılacak\n- [x] tamamlandı")

    texts = _paragraph_texts(docx.Document(io.BytesIO(data)))

    assert any("☐" in t and "yapılacak" in t for t in texts)
    assert any("☑" in t and "tamamlandı" in t for t in texts)


def test_callout_renders_type_label_and_body():
    data = render_docx(None, "> [!WARNING]\n> dikkatli olun")

    texts = _paragraph_texts(docx.Document(io.BytesIO(data)))

    assert any("WARNING" in t and "dikkatli olun" in t for t in texts)


def test_table_cell_preserves_bold_and_column_alignment():
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    markdown = "| Ad |\n|---:|\n| **önemli** |"
    data = render_docx(None, markdown)

    document = docx.Document(io.BytesIO(data))
    cell = document.tables[0].rows[1].cells[0]

    assert cell.paragraphs[0].runs[0].bold is True
    assert cell.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT


def test_image_from_data_uri_is_embedded():
    data = render_docx(None, f"![alt metni]({_PNG_DATA_URI})")

    document = docx.Document(io.BytesIO(data))

    assert len(document.inline_shapes) == 1


def test_image_with_unloadable_source_does_not_break_document():
    data = render_docx(None, "![alt](https://does-not-exist.invalid/x.png)\n\nsonrası metin")

    document = docx.Document(io.BytesIO(data))

    assert "sonrası metin" in _paragraph_texts(document)


def test_horizontal_rule_adds_bottom_border_paragraph():
    data = render_docx(None, "bir\n\n---\n\niki")

    document = docx.Document(io.BytesIO(data))
    borders = document.element.body.findall(".//" + qn("w:pBdr"))

    assert len(borders) == 1


def test_page_break_marker_inserts_word_page_break():
    data = render_docx(None, "bir\n\n<!-- pagebreak -->\n\niki")

    document = docx.Document(io.BytesIO(data))
    breaks = document.element.body.findall(".//" + qn("w:br"))

    assert any(b.get(qn("w:type")) == "page" for b in breaks)


def test_footnote_definition_renders_marker_and_text():
    markdown = "iddia[^1]\n\n[^1]: kaynak açıklaması"
    data = render_docx(None, markdown)

    texts = _paragraph_texts(docx.Document(io.BytesIO(data)))

    assert any("[1]" in t and "kaynak açıklaması" in t for t in texts)


def test_table_caption_renders_below_table():
    markdown = "| A |\n| --- |\n| 1 |\n*Tablo: Örnek veri*"
    data = render_docx(None, markdown)

    texts = _paragraph_texts(docx.Document(io.BytesIO(data)))

    assert any("Örnek veri" in t for t in texts)
