"""Tests for the DOCX cover page: title/subtitle/metadata stack, logo, and
the page break that separates it from the main content."""

import io

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from service.documents.docx_renderer import render_docx
from service.documents.options import parse_document_options

_PNG_DATA_URI = (
    "data:image/png;base64,"
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
    "+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _page_break_count(document) -> int:
    breaks = document.element.body.findall(f".//{qn('w:br')}")
    return sum(1 for b in breaks if b.get(qn("w:type")) == "page")


def test_cover_disabled_by_default_no_extra_page_break():
    data = render_docx("Rapor", "içerik")

    document = docx.Document(io.BytesIO(data))
    assert _page_break_count(document) == 0
    assert document.paragraphs[0].text == "Rapor"


def test_cover_page_shows_title_subtitle_and_metadata_centered():
    options = parse_document_options(
        {
            "cover": {
                "enabled": True,
                "subtitle": "Yazılım Gereksinim Spesifikasyonu",
                "project": "AgenticAI",
                "version": "1.0",
                "date": "2026-08-11",
                "author": "Ali Burak Pekışık",
                "organization": "Acme",
            }
        }
    )

    data = render_docx("SRS", "içerik", options=options)

    document = docx.Document(io.BytesIO(data))
    texts = [p.text for p in document.paragraphs]

    assert any(
        p.text == "SRS" and p.alignment == WD_ALIGN_PARAGRAPH.CENTER for p in document.paragraphs
    )
    assert any("Yazılım Gereksinim Spesifikasyonu" in t for t in texts)
    assert any("AgenticAI" in t for t in texts)
    assert any("1.0" in t for t in texts)
    assert any("2026-08-11" in t for t in texts)
    assert any("Ali Burak Pekışık" in t for t in texts)
    assert any("Acme" in t for t in texts)


def test_cover_page_ends_with_a_page_break_before_content():
    options = parse_document_options({"cover": {"enabled": True}})

    data = render_docx("SRS", "# Giriş", options=options)

    document = docx.Document(io.BytesIO(data))
    assert _page_break_count(document) >= 1


def test_cover_page_classification_is_bold():
    options = parse_document_options({"cover": {"enabled": True, "classification": "GİZLİ"}})

    data = render_docx(None, "x", options=options)

    document = docx.Document(io.BytesIO(data))
    classification_paragraph = next(p for p in document.paragraphs if p.text == "GİZLİ")
    assert classification_paragraph.runs[0].bold is True


def test_cover_page_embeds_logo_image():
    options = parse_document_options({"cover": {"enabled": True, "logo": _PNG_DATA_URI}})

    data = render_docx("SRS", "x", options=options)

    document = docx.Document(io.BytesIO(data))
    assert len(document.inline_shapes) == 1
