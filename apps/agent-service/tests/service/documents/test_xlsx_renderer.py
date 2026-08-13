"""Tests for XLSX styling enrichment (Group E): themed header row, freeze
pane, autofilter, zebra striping, column widths/formats/alignments and
static conditional formatting."""

import io

import openpyxl

from service.documents.options import parse_spreadsheet_options
from service.documents.themes import resolve_theme
from service.documents.xlsx_renderer import render_xlsx


def _load(data: bytes):
    return openpyxl.load_workbook(io.BytesIO(data))


def _fill_matches(cell, hex_color: str) -> bool:
    rgb = cell.fill.start_color.rgb
    return bool(rgb) and rgb.upper().endswith(hex_color.lstrip("#").upper())


_ROWS = [["Ad", "Durum"], ["Bir", "PASS"], ["İki", "FAIL"], ["Üç", "PASS"]]


def test_header_row_gets_bold_font_and_themed_fill_by_default():
    data = render_xlsx([{"name": "S", "rows": _ROWS}])

    ws = _load(data)["S"]
    default_theme = resolve_theme("default")
    assert ws["A1"].font.bold is True
    assert _fill_matches(ws["A1"], default_theme.table_header_bg)


def test_named_theme_changes_header_fill_color():
    options = parse_spreadsheet_options({"theme": "dark_accent"})

    data = render_xlsx([{"name": "S", "rows": _ROWS}], options=options)

    ws = _load(data)["S"]
    dark_accent = resolve_theme("dark_accent")
    assert _fill_matches(ws["A1"], dark_accent.table_header_bg)


def test_freeze_header_freezes_row_two_by_default():
    data = render_xlsx([{"name": "S", "rows": _ROWS}])

    assert _load(data)["S"].freeze_panes == "A2"


def test_freeze_header_disabled_via_options():
    options = parse_spreadsheet_options({"freeze_header": False})

    data = render_xlsx([{"name": "S", "rows": _ROWS}], options=options)

    assert _load(data)["S"].freeze_panes is None


def test_autofilter_applied_by_default():
    data = render_xlsx([{"name": "S", "rows": _ROWS}])

    assert _load(data)["S"].auto_filter.ref == "A1:B4"


def test_autofilter_disabled_via_options():
    options = parse_spreadsheet_options({"autofilter": False})

    data = render_xlsx([{"name": "S", "rows": _ROWS}], options=options)

    assert not _load(data)["S"].auto_filter.ref


def test_zebra_shades_alternating_data_rows():
    options = parse_spreadsheet_options({"zebra": True})

    data = render_xlsx([{"name": "S", "rows": _ROWS}], options=options)

    ws = _load(data)["S"]
    default_theme = resolve_theme("default")
    assert not _fill_matches(ws["A2"], default_theme.table_zebra_bg)
    assert _fill_matches(ws["A3"], default_theme.table_zebra_bg)


def test_column_widths_applied():
    options = parse_spreadsheet_options({"column_widths": [30, 12]})

    data = render_xlsx([{"name": "S", "rows": _ROWS}], options=options)

    ws = _load(data)["S"]
    assert ws.column_dimensions["A"].width == 30
    assert ws.column_dimensions["B"].width == 12


def test_column_formats_apply_number_format():
    rows = [["Ürün", "Oran"], ["A", 0.5]]
    options = parse_spreadsheet_options({"column_formats": ["text", "percent"]})

    data = render_xlsx([{"name": "S", "rows": rows}], options=options)

    ws = _load(data)["S"]
    assert ws["B2"].number_format == "0%"


def test_column_alignments_applied():
    options = parse_spreadsheet_options({"column_alignments": ["left", "center"]})

    data = render_xlsx([{"name": "S", "rows": _ROWS}], options=options)

    ws = _load(data)["S"]
    assert ws["B2"].alignment.horizontal == "center"


def test_conditional_format_shades_only_matching_cells():
    options = parse_spreadsheet_options(
        {
            "conditional_formats": [
                {"column": 2, "rule": "equals", "value": "FAIL", "color": "#FFC7CE"}
            ]
        }
    )

    data = render_xlsx([{"name": "S", "rows": _ROWS}], options=options)

    ws = _load(data)["S"]
    assert _fill_matches(ws["B3"], "#FFC7CE")
    assert not _fill_matches(ws["B2"], "#FFC7CE")


def test_autofit_columns_widens_column_for_long_text():
    rows = [["Kısa"], ["Bu hücre oldukça uzun bir metin içeriyor ve genişlik almalı"]]

    data = render_xlsx([{"name": "S", "rows": rows}])

    ws = _load(data)["S"]
    assert ws.column_dimensions["A"].width > 15
