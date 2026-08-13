"""Tests for DOCX page scaffolding (size/orientation/margins) and theme
application (fonts/colors on headings, body text, tables, code blocks)."""

import io

import docx
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, RGBColor

from service.documents.docx_renderer import render_docx
from service.documents.options import parse_document_options
from service.documents.themes import resolve_theme

# Word stores lengths in twips (1/1440in); a millimeter value that isn't a
# whole number of twips round-trips with a sub-twip rounding error, so page
# geometry assertions compare within one twip (635 EMU) rather than exactly.
_TWIP_EMU = 635


def _assert_mm_close(actual, expected_mm) -> None:
    assert abs(actual - Mm(expected_mm)) <= _TWIP_EMU


# ---------------------------------------------------------------------------
# Page setup
# ---------------------------------------------------------------------------


def test_default_options_produce_a4_portrait():
    data = render_docx(None, "x")

    section = docx.Document(io.BytesIO(data)).sections[0]

    assert section.orientation == WD_ORIENT.PORTRAIT
    _assert_mm_close(section.page_width, 210)
    _assert_mm_close(section.page_height, 297)


def test_letter_landscape_swaps_dimensions():
    options = parse_document_options({"page": {"size": "Letter", "orientation": "landscape"}})

    data = render_docx(None, "x", options=options)
    section = docx.Document(io.BytesIO(data)).sections[0]

    assert section.orientation == WD_ORIENT.LANDSCAPE
    _assert_mm_close(section.page_width, 279.4)
    _assert_mm_close(section.page_height, 215.9)


def test_custom_margins_are_applied():
    options = parse_document_options(
        {"page": {"margins": {"top": 10, "bottom": 12, "left": 15, "right": 18}}}
    )

    data = render_docx(None, "x", options=options)
    section = docx.Document(io.BytesIO(data)).sections[0]

    _assert_mm_close(section.top_margin, 10)
    _assert_mm_close(section.bottom_margin, 12)
    _assert_mm_close(section.left_margin, 15)
    _assert_mm_close(section.right_margin, 18)


# ---------------------------------------------------------------------------
# Theme — headings, body, links
# ---------------------------------------------------------------------------


def test_default_options_still_apply_the_default_theme():
    """§13.1: even with no `options` argument, the default theme's font and
    heading color are applied — the tool call never regresses to bare Word
    defaults."""
    data = render_docx(None, "# Başlık\n\ngövde metni")

    document = docx.Document(io.BytesIO(data))
    heading_run = document.paragraphs[0].runs[0]
    body_run = document.paragraphs[1].runs[0]
    default_theme = resolve_theme("default")

    assert heading_run.font.name == default_theme.heading_font
    assert heading_run.font.color.rgb == RGBColor.from_string(
        default_theme.heading_color.lstrip("#")
    )
    assert body_run.font.name == default_theme.body_font


def test_named_theme_changes_heading_font_and_color():
    options = parse_document_options({"theme": "academic"})

    data = render_docx(None, "# Başlık", options=options)

    heading_run = docx.Document(io.BytesIO(data)).paragraphs[0].runs[0]
    academic = resolve_theme("academic")

    assert heading_run.font.name == academic.heading_font
    assert heading_run.font.color.rgb == RGBColor.from_string(academic.heading_color.lstrip("#"))


def test_explicit_font_override_wins_over_theme():
    options = parse_document_options({"theme": "academic", "font": {"body": "Arial"}})

    data = render_docx(None, "gövde", options=options)

    body_run = docx.Document(io.BytesIO(data)).paragraphs[0].runs[0]

    assert body_run.font.name == "Arial"


def test_link_uses_theme_link_color():
    options = parse_document_options({"theme": "dark_accent"})

    data = render_docx(None, "[metin](https://example.com)", options=options)

    from docx.oxml.ns import qn

    document = docx.Document(io.BytesIO(data))
    color_el = document.element.body.find(f".//{qn('w:hyperlink')}//{qn('w:color')}")
    dark_accent = resolve_theme("dark_accent")

    assert color_el.get(qn("w:val")).upper() == dark_accent.link_color.lstrip("#").upper()


# ---------------------------------------------------------------------------
# Theme — tables and code blocks
# ---------------------------------------------------------------------------


def _cell_shading_fill(cell) -> str | None:
    from docx.oxml.ns import qn

    tc_pr = cell._tc.find(qn("w:tcPr"))
    if tc_pr is None:
        return None
    shd = tc_pr.find(qn("w:shd"))
    return shd.get(qn("w:fill")) if shd is not None else None


def test_table_header_row_is_shaded_and_bold():
    data = render_docx(None, "| A |\n| --- |\n| 1 |")

    document = docx.Document(io.BytesIO(data))
    table = document.tables[0]
    header_cell = table.rows[0].cells[0]
    default_theme = resolve_theme("default")

    assert header_cell.paragraphs[0].runs[0].bold is True
    assert (
        _cell_shading_fill(header_cell).upper() == default_theme.table_header_bg.lstrip("#").upper()
    )


def test_zebra_table_style_shades_alternating_data_rows():
    options = parse_document_options({"tables": {"style": "zebra"}})
    markdown = "| A |\n| --- |\n| 1 |\n| 2 |\n| 3 |"

    data = render_docx(None, markdown, options=options)

    table = docx.Document(io.BytesIO(data)).tables[0]
    assert _cell_shading_fill(table.rows[1].cells[0]) is None
    assert _cell_shading_fill(table.rows[2].cells[0]) is not None


def test_code_block_paragraph_is_shaded_with_theme_code_background():
    from docx.oxml.ns import qn

    data = render_docx(None, "```\nx = 1\n```")

    document = docx.Document(io.BytesIO(data))
    code_paragraph = next(p for p in document.paragraphs if p.text == "x = 1")
    p_pr = code_paragraph._p.find(qn("w:pPr"))
    shd = p_pr.find(qn("w:shd"))
    default_theme = resolve_theme("default")

    assert shd.get(qn("w:fill")).upper() == default_theme.code_bg.lstrip("#").upper()


def test_table_alignment_still_applied_alongside_theme_shading():
    data = render_docx(None, "| A |\n|---:|\n| 1 |")

    document = docx.Document(io.BytesIO(data))
    assert document.tables[0].rows[0].cells[0].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
