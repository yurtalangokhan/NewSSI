"""Tests for the always-on document output tools (create_document, create_spreadsheet)."""

from __future__ import annotations

import json
from typing import Any
from unittest.mock import AsyncMock

import pytest
from langchain_core.runnables import RunnableConfig

from agents import document_tools


def _config(user_id: str = "user-1", thread_id: str = "thread-1") -> RunnableConfig:
    return RunnableConfig(configurable={"user_id": user_id, "thread_id": thread_id})


class _FakeDocumentRepository:
    def __init__(self) -> None:
        self.created: list[dict[str, Any]] = []

    async def create(self, **kwargs: Any) -> dict[str, Any]:
        self.created.append(kwargs)
        return {"id": "doc-1", **kwargs}


@pytest.fixture(autouse=True)
def _clear_file_store():
    from service.FileService import _STORE

    _STORE.clear()
    yield
    _STORE.clear()


@pytest.fixture
def fake_repo(monkeypatch) -> _FakeDocumentRepository:
    repo = _FakeDocumentRepository()
    monkeypatch.setattr("core.db.repositories.document_repo.DocumentRepository", lambda: repo)
    return repo


@pytest.fixture
def fake_minio_upload(monkeypatch):
    calls: list[dict[str, Any]] = []

    def _upload(*, user_id, file_id, filename, data, mime_type):
        calls.append(
            {
                "user_id": user_id,
                "file_id": file_id,
                "filename": filename,
                "data": data,
                "mime_type": mime_type,
            }
        )
        return f"documents/{user_id}/{file_id}/{filename}"

    monkeypatch.setattr("service.MinioService.upload_file", _upload)
    return calls


# ---------------------------------------------------------------------------
# create_document
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_create_document_returns_generated_file_payload(fake_repo, fake_minio_upload):
    result = await document_tools.create_document.ainvoke(
        {
            "filename": "rapor",
            "format": "pdf",
            "content": "# Başlık\n\nİçerik",
        },
        _config(),
    )

    payload = json.loads(result)
    assert payload["__generated_file__"] is True
    assert payload["filename"] == "rapor.pdf"
    assert payload["mime_type"] == "application/pdf"
    assert payload["size_bytes"] > 0
    assert payload["download_url"] == f"/api/chat/file/{payload['file_id']}?download=1"


@pytest.mark.asyncio
async def test_create_document_stores_bytes_in_file_service(fake_repo, fake_minio_upload):
    from service.FileService import get_file

    result = await document_tools.create_document.ainvoke(
        {
            "filename": "rapor",
            "format": "txt",
            "content": "merhaba dünya",
        },
        _config(),
    )

    payload = json.loads(result)
    record = get_file(payload["file_id"])
    assert record is not None
    assert record.data == b"merhaba d\xc3\xbcnya"
    assert record.filename == "rapor.txt"


@pytest.mark.asyncio
async def test_create_document_uploads_to_minio_with_user_and_file_id(fake_repo, fake_minio_upload):
    result = await document_tools.create_document.ainvoke(
        {
            "filename": "rapor",
            "format": "md",
            "content": "gövde",
        },
        _config(user_id="user-42"),
    )

    payload = json.loads(result)
    assert len(fake_minio_upload) == 1
    assert fake_minio_upload[0]["user_id"] == "user-42"
    assert fake_minio_upload[0]["file_id"] == payload["file_id"]
    assert fake_minio_upload[0]["filename"] == "rapor.md"


@pytest.mark.asyncio
async def test_create_document_persists_document_row_scoped_to_thread(fake_repo, fake_minio_upload):
    await document_tools.create_document.ainvoke(
        {
            "filename": "rapor",
            "format": "md",
            "content": "gövde",
        },
        _config(user_id="user-42", thread_id="thread-99"),
    )

    assert len(fake_repo.created) == 1
    assert fake_repo.created[0]["user_id"] == "user-42"
    assert fake_repo.created[0]["thread_id"] == "thread-99"
    assert fake_repo.created[0]["filename"] == "rapor.md"


@pytest.mark.asyncio
async def test_create_document_rejects_unsupported_format(fake_repo, fake_minio_upload):
    result = await document_tools.create_document.ainvoke(
        {
            "filename": "rapor",
            "format": "rtf",
            "content": "gövde",
        },
        _config(),
    )

    assert "__generated_file__" not in result
    assert "rtf" in result.lower() or "format" in result.lower()
    assert fake_repo.created == []
    assert fake_minio_upload == []


@pytest.mark.asyncio
async def test_create_document_rejects_empty_content(fake_repo, fake_minio_upload):
    result = await document_tools.create_document.ainvoke(
        {
            "filename": "rapor",
            "format": "pdf",
            "content": "   ",
        },
        _config(),
    )

    assert "__generated_file__" not in result
    assert fake_repo.created == []


@pytest.mark.asyncio
async def test_create_document_survives_minio_upload_failure(fake_repo, monkeypatch):
    def _boom(**kwargs):
        raise RuntimeError("minio unreachable")

    monkeypatch.setattr("service.MinioService.upload_file", _boom)

    result = await document_tools.create_document.ainvoke(
        {
            "filename": "rapor",
            "format": "txt",
            "content": "gövde",
        },
        _config(),
    )

    payload = json.loads(result)
    assert payload["__generated_file__"] is True


@pytest.mark.asyncio
async def test_create_document_survives_db_persist_failure(fake_minio_upload, monkeypatch):
    class _BoomRepo:
        async def create(self, **kwargs):
            raise RuntimeError("db unavailable")

    monkeypatch.setattr(
        "core.db.repositories.document_repo.DocumentRepository", lambda: _BoomRepo()
    )

    result = await document_tools.create_document.ainvoke(
        {
            "filename": "rapor",
            "format": "txt",
            "content": "gövde",
        },
        _config(),
    )

    payload = json.loads(result)
    assert payload["__generated_file__"] is True


# ---------------------------------------------------------------------------
# create_document — options
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_create_document_applies_theme_from_options(fake_repo, fake_minio_upload):
    import io

    import docx

    from service.documents.themes import resolve_theme
    from service.FileService import get_file

    result = await document_tools.create_document.ainvoke(
        {
            "filename": "rapor",
            "format": "docx",
            "content": "# Başlık",
            "options": {"theme": "academic"},
        },
        _config(),
    )

    payload = json.loads(result)
    record = get_file(payload["file_id"])
    document = docx.Document(io.BytesIO(record.data))
    academic = resolve_theme("academic")
    assert document.paragraphs[0].runs[0].font.name == academic.heading_font


@pytest.mark.asyncio
async def test_create_document_accepts_options_as_json_string(fake_repo, fake_minio_upload):
    result = await document_tools.create_document.ainvoke(
        {
            "filename": "rapor",
            "format": "pdf",
            "content": "gövde",
            "options": '{"theme": "corporate_blue"}',
        },
        _config(),
    )

    payload = json.loads(result)
    assert payload["__generated_file__"] is True


@pytest.mark.asyncio
async def test_create_document_never_fails_on_malformed_options(fake_repo, fake_minio_upload):
    """A single bad options key must not cost the model a full regeneration —
    the tool falls back to defaults for whatever it cannot make sense of."""
    result = await document_tools.create_document.ainvoke(
        {
            "filename": "rapor",
            "format": "pdf",
            "content": "gövde",
            "options": {"font": {"size": 9999}, "theme": "does_not_exist", "toc": "not-an-object"},
        },
        _config(),
    )

    payload = json.loads(result)
    assert payload["__generated_file__"] is True


# ---------------------------------------------------------------------------
# create_spreadsheet
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_create_spreadsheet_returns_generated_file_payload(fake_repo, fake_minio_upload):
    result = await document_tools.create_spreadsheet.ainvoke(
        {
            "filename": "satis",
            "format": "xlsx",
            "sheets": [{"name": "Q1", "rows": [["Ürün", "Adet"], ["Çilek", 3]]}],
        },
        _config(),
    )

    payload = json.loads(result)
    assert payload["__generated_file__"] is True
    assert payload["filename"] == "satis.xlsx"
    assert (
        payload["mime_type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@pytest.mark.asyncio
async def test_create_spreadsheet_csv_round_trips_rows(fake_repo, fake_minio_upload):
    from service.FileService import get_file

    result = await document_tools.create_spreadsheet.ainvoke(
        {
            "filename": "satis",
            "format": "csv",
            "sheets": [{"name": "Q1", "rows": [["a", "b"], ["1", "2"]]}],
        },
        _config(),
    )

    payload = json.loads(result)
    record = get_file(payload["file_id"])
    assert record.data.decode("utf-8-sig").splitlines() == ["a,b", "1,2"]


@pytest.mark.asyncio
async def test_create_spreadsheet_rejects_unsupported_format(fake_repo, fake_minio_upload):
    result = await document_tools.create_spreadsheet.ainvoke(
        {
            "filename": "satis",
            "format": "pdf",
            "sheets": [{"name": "Q1", "rows": [["a"]]}],
        },
        _config(),
    )

    assert "__generated_file__" not in result
    assert fake_repo.created == []


@pytest.mark.asyncio
async def test_create_spreadsheet_rejects_no_sheets(fake_repo, fake_minio_upload):
    result = await document_tools.create_spreadsheet.ainvoke(
        {
            "filename": "satis",
            "format": "xlsx",
            "sheets": [],
        },
        _config(),
    )

    assert "__generated_file__" not in result
    assert fake_repo.created == []


@pytest.mark.asyncio
async def test_create_spreadsheet_applies_zebra_from_options(fake_repo, fake_minio_upload):
    import io

    import openpyxl

    from service.FileService import get_file

    result = await document_tools.create_spreadsheet.ainvoke(
        {
            "filename": "satis",
            "format": "xlsx",
            "sheets": [{"name": "Q1", "rows": [["Ürün"], ["a"], ["b"], ["c"]]}],
            "options": {"zebra": True},
        },
        _config(),
    )

    payload = json.loads(result)
    record = get_file(payload["file_id"])
    worksheet = openpyxl.load_workbook(io.BytesIO(record.data))["Q1"]
    assert worksheet["A3"].fill.start_color.rgb not in (None, "00000000")


@pytest.mark.asyncio
async def test_create_spreadsheet_never_fails_on_malformed_options(fake_repo, fake_minio_upload):
    result = await document_tools.create_spreadsheet.ainvoke(
        {
            "filename": "satis",
            "format": "xlsx",
            "sheets": [{"name": "Q1", "rows": [["a"]]}],
            "options": {"theme": "does_not_exist", "column_widths": "not-a-list"},
        },
        _config(),
    )

    payload = json.loads(result)
    assert payload["__generated_file__"] is True


# ---------------------------------------------------------------------------
# get_document_tools / bind_document_tools
# ---------------------------------------------------------------------------


def test_get_document_tools_returns_both_tools_when_enabled(monkeypatch):
    monkeypatch.setattr(document_tools.settings, "DOCUMENT_TOOLS_ENABLED", True)

    tools = document_tools.get_document_tools()

    assert {t.name for t in tools} == {"create_document", "create_spreadsheet"}


def test_get_document_tools_returns_empty_when_disabled(monkeypatch):
    monkeypatch.setattr(document_tools.settings, "DOCUMENT_TOOLS_ENABLED", False)

    assert document_tools.get_document_tools() == []


def test_get_document_tools_uses_rich_description_by_default(monkeypatch):
    monkeypatch.setattr(document_tools.settings, "DOCUMENT_TOOLS_ENABLED", True)
    monkeypatch.setattr(document_tools.settings, "DOCUMENT_TOOLS_RICH_OPTIONS", True)

    tools = {t.name: t for t in document_tools.get_document_tools()}

    assert "options" in tools["create_document"].description
    assert "options" in tools["create_spreadsheet"].description


def test_get_document_tools_uses_short_description_when_rich_options_disabled(monkeypatch):
    monkeypatch.setattr(document_tools.settings, "DOCUMENT_TOOLS_ENABLED", True)
    monkeypatch.setattr(document_tools.settings, "DOCUMENT_TOOLS_RICH_OPTIONS", False)

    tools = {t.name: t for t in document_tools.get_document_tools()}

    assert len(tools["create_document"].description) < len(
        document_tools.create_document.description
    )


def test_bind_document_tools_returns_bound_model_and_true_on_success():
    fake_model = AsyncMock()
    fake_model.bind_tools = lambda tools: "bound-model"

    bound, supported = document_tools.bind_document_tools(fake_model, [])

    assert bound == "bound-model"
    assert supported is True


def test_bind_document_tools_falls_back_when_bind_tools_unsupported():
    class NoToolSupport:
        def bind_tools(self, tools):
            raise NotImplementedError("this model does not support tools")

    model = NoToolSupport()

    bound, supported = document_tools.bind_document_tools(model, [])

    assert bound is model
    assert supported is False


# ---------------------------------------------------------------------------
# DOCUMENT_TOOL_PROMPT — must explicitly steer weak models away from calling
# the tools on greetings/small talk (regression: a local model created a
# document titled "merhaba" in response to a plain greeting).
# ---------------------------------------------------------------------------


def test_document_tool_prompt_explicitly_forbids_calling_on_greetings():
    prompt = document_tools.DOCUMENT_TOOL_PROMPT.lower()

    assert "greeting" in prompt or "small talk" in prompt
    assert "do not" in prompt or "don't" in prompt
    assert "when in doubt" in prompt


def test_document_tool_prompt_mentions_options():
    assert "options" in document_tools.DOCUMENT_TOOL_PROMPT.lower()


def test_document_tool_prompt_forbids_repeating_the_download_link():
    """The UI already renders a clickable download card from the tool
    result; a second link in the model's own reply text opens a preview
    modal that doesn't know the file's real type and fails."""
    prompt = document_tools.DOCUMENT_TOOL_PROMPT.lower()

    assert "download_url" in prompt
    assert "do not include" in prompt or "don't include" in prompt


class TestRecoverDocumentToolArgs:
    """Local models often put the document body somewhere other than `content`.

    Rejecting those calls costs a full regeneration of a long document, so the
    body is recovered wherever it is unambiguous.
    """

    def test_leaves_a_well_formed_call_untouched(self):
        args = {"filename": "rapor", "format": "pdf", "content": "# Başlık"}

        assert document_tools.recover_document_tool_args(args) == args

    def test_recovers_the_body_from_an_alias_key(self):
        recovered = document_tools.recover_document_tool_args(
            {"filename": "rapor", "format": "pdf", "text": "# Başlık\n\nGövde"}
        )

        assert recovered["content"] == "# Başlık\n\nGövde"
        assert "text" not in recovered

    def test_recovers_a_body_wrapped_in_pseudo_xml_inside_another_argument(self):
        recovered = document_tools.recover_document_tool_args(
            {
                "filename": "rapor",
                "format": "docx",
                "title": "<content># Başlık\n\nGövde</content>",
            }
        )

        assert recovered["content"] == "# Başlık\n\nGövde"
        # The wrapper was the whole value, so the borrowed argument is dropped.
        assert "title" not in recovered

    def test_keeps_the_rest_of_a_partially_wrapped_argument(self):
        recovered = document_tools.recover_document_tool_args(
            {
                "filename": "rapor",
                "format": "docx",
                "title": "Rapor <content>Gövde</content>",
            }
        )

        assert recovered["content"] == "Gövde"
        assert recovered["title"] == "Rapor"

    def test_does_not_invent_content_from_an_unrelated_argument(self):
        """A description is not the document body — let validation reject it."""
        args = {
            "filename": "rapor",
            "format": "pdf",
            "description": "Bu dosya üniversite tarihi içerir.",
        }

        assert document_tools.recover_document_tool_args(args) == args

    def test_ignores_blank_aliases(self):
        args = {"filename": "rapor", "format": "pdf", "body": "   "}

        assert document_tools.recover_document_tool_args(args) == args

    def test_recovers_a_body_that_only_kept_its_closing_tag(self):
        """The provider's own parser often eats the opening tag."""
        recovered = document_tools.recover_document_tool_args(
            {
                "filename": "hacet_tarihcesi.docx",
                "format": "docx",
                "belge": "# Hacettepe\n\n" + "Uzun gövde. " * 30 + "\n(Türkçe)</content>",
            }
        )

        assert recovered["content"].startswith("# Hacettepe")
        assert "</content>" not in recovered["content"]
        assert recovered["filename"] == "hacet_tarihcesi.docx"

    def test_recovers_a_long_multiline_stray_argument(self):
        body = "# Başlık\n\n" + "Gövde metni. " * 40
        recovered = document_tools.recover_document_tool_args(
            {"filename": "rapor", "format": "pdf", "icerik": body}
        )

        assert recovered["content"] == body.strip()

    def test_cleans_markup_out_of_a_content_that_was_passed_correctly(self):
        recovered = document_tools.recover_document_tool_args(
            {
                "filename": "rapor",
                "format": "pdf",
                "content": "<content># Başlık\n\nGövde</content>",
            }
        )

        assert recovered["content"] == "# Başlık\n\nGövde"

    def test_never_promotes_a_short_single_line_argument(self):
        args = {"filename": "rapor", "format": "pdf", "konu": "üniversite tarihi"}

        assert document_tools.recover_document_tool_args(args) == args

    def test_never_borrows_the_filename_as_the_body(self):
        args = {"filename": "x" * 400 + "\nsatır", "format": "pdf"}

        assert document_tools.recover_document_tool_args(args) == args


class TestRecoverDocumentToolArgsOptions:
    """Some models flatten style keys onto the top-level call instead of
    nesting them under `options` — these are gathered rather than rejected."""

    def test_gathers_a_stray_top_level_theme_into_options(self):
        recovered = document_tools.recover_document_tool_args(
            {"filename": "rapor", "format": "docx", "content": "x", "theme": "corporate_blue"}
        )

        assert recovered["options"] == {"theme": "corporate_blue"}
        assert "theme" not in recovered

    def test_gathers_multiple_stray_keys(self):
        recovered = document_tools.recover_document_tool_args(
            {
                "filename": "rapor",
                "format": "docx",
                "content": "x",
                "toc": {"enabled": True},
                "cover": {"enabled": True},
            }
        )

        assert recovered["options"] == {"toc": {"enabled": True}, "cover": {"enabled": True}}

    def test_merges_stray_keys_with_an_existing_options_dict(self):
        recovered = document_tools.recover_document_tool_args(
            {
                "filename": "rapor",
                "format": "docx",
                "content": "x",
                "options": {"toc": {"enabled": True}},
                "theme": "academic",
            }
        )

        assert recovered["options"] == {"toc": {"enabled": True}, "theme": "academic"}

    def test_explicit_options_value_wins_over_a_stray_duplicate(self):
        recovered = document_tools.recover_document_tool_args(
            {
                "filename": "rapor",
                "format": "docx",
                "content": "x",
                "options": {"theme": "academic"},
                "theme": "corporate_blue",
            }
        )

        assert recovered["options"] == {"theme": "academic"}

    def test_leaves_args_untouched_when_no_stray_option_keys_present(self):
        args = {"filename": "rapor", "format": "docx", "content": "x"}

        assert document_tools.recover_document_tool_args(args) == args
